import json, os, re, sys, traceback
from pathlib import Path
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from io import BytesIO
OUT=Path("research/out"); OUT.mkdir(parents=True, exist_ok=True)
UA={"User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/154.0 Safari/537.36","Accept":"text/html,application/xhtml+xml,application/json,*/*"}

def dump(name,obj):
    (OUT/name).write_text(json.dumps(obj,indent=2,default=str),encoding="utf-8")

# PHOENIX
try:
    url="https://maps.phoenix.gov/pub/rest/services/Public/Planning_Permit/MapServer/1/query"
    params={"where":"PER_ENT_DATE IS NOT NULL","outFields":"*","returnGeometry":"false","orderByFields":"PER_ENT_DATE DESC,OBJECTID DESC","resultRecordCount":"1000","f":"json"}
    r=requests.get(url,params=params,headers=UA,timeout=60); r.raise_for_status()
    j=r.json(); dump("phoenix_1000.json",j)
    feats=[x.get("attributes",{}) for x in j.get("features",[])]
    dist={}
    for fld in ["PER_TYPE","PER_TYPE_DESC","MOD_DESC","SCOPE_CODE","PERMIT_STAT"]:
        d={}
        for a in feats:
            k=str(a.get(fld) or "")
            d[k]=d.get(k,0)+1
        dist[fld]=sorted(d.items(),key=lambda kv:(-kv[1],kv[0]))[:100]
    dump("phoenix_diag.json",{"status":r.status_code,"url":r.url,"count":len(feats),"fields":list(feats[0]) if feats else [],"distributions":dist,"sample":feats[:25]})
except Exception as e:
    dump("phoenix_error.json",{"error":repr(e),"trace":traceback.format_exc()})

# DALLAS CO
try:
    co="https://dallascityhall.com/departments/pnv/Documents/AH%20Memos/COSMonthly%20Report2026-Aug.xlsx"
    r=requests.get(co,headers=UA,timeout=60); 
    dump("dallas_co_http.json",{"status":r.status_code,"url":r.url,"ctype":r.headers.get("content-type"),"bytes":len(r.content)})
    r.raise_for_status()
    (OUT/"dallas_aug.xlsx").write_bytes(r.content)
    wb=load_workbook(BytesIO(r.content),read_only=True,data_only=True)
    info={}
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        info[ws.title]={"max_row":ws.max_row,"max_col":ws.max_column,"rows":[list(x) for x in rows[:30]]}
    dump("dallas_co_diag.json",info)
except Exception as e:
    dump("dallas_co_error.json",{"error":repr(e),"trace":traceback.format_exc()})

# TDLR form diagnostics
try:
    u="https://www.tdlr.texas.gov/TABS/Search/"
    s=requests.Session(); r=s.get(u,headers=UA,timeout=60); r.raise_for_status()
    (OUT/"tdlr_search.html").write_text(r.text,encoding="utf-8")
    soup=BeautifulSoup(r.text,"lxml")
    forms=[]
    for f in soup.find_all("form"):
        fields=[]
        for el in f.find_all(["input","select","button","textarea"]):
            fields.append({"tag":el.name,"name":el.get("name"),"id":el.get("id"),"type":el.get("type"),"value":el.get("value"),"text":el.get_text(" ",strip=True)[:200]})
        forms.append({"action":f.get("action"),"method":f.get("method"),"id":f.get("id"),"fields":fields})
    scripts=[x.get("src") for x in soup.find_all("script") if x.get("src")]
    dump("tdlr_diag.json",{"status":r.status_code,"url":r.url,"forms":forms,"scripts":scripts})
except Exception as e:
    dump("tdlr_error.json",{"error":repr(e),"trace":traceback.format_exc()})

# ATLANTA ACCELA diagnostics
try:
    u="https://aca-prod.accela.com/ATLANTA_GA/Cap/CapHome.aspx?module=Building&TabName=Building"
    s=requests.Session(); r=s.get(u,headers=UA,timeout=60,allow_redirects=True)
    (OUT/"atlanta_cap.html").write_text(r.text,encoding="utf-8")
    soup=BeautifulSoup(r.text,"lxml")
    forms=[]
    for f in soup.find_all("form"):
        fields=[]
        for el in f.find_all(["input","select","button","textarea","a"]):
            name=el.get("name"); eid=el.get("id")
            if name or eid or el.name=="a":
                fields.append({"tag":el.name,"name":name,"id":eid,"type":el.get("type"),"value":el.get("value"),"href":el.get("href"),"text":el.get_text(" ",strip=True)[:200]})
        forms.append({"action":f.get("action"),"method":f.get("method"),"id":f.get("id"),"fields":fields})
    dump("atlanta_diag.json",{"status":r.status_code,"url":r.url,"ctype":r.headers.get("content-type"),"bytes":len(r.content),"title":soup.title.string if soup.title else None,"forms":forms[:3],"scripts":[x.get("src") for x in soup.find_all("script") if x.get("src")]})
except Exception as e:
    dump("atlanta_error.json",{"error":repr(e),"trace":traceback.format_exc()})
